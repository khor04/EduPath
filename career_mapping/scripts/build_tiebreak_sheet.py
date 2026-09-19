"""Blind tie-break workbook for a third reviewer: only titles where Reviewers A and B
picked no common option. Same options/letters as the originals (original 'No' kept so
answers map back through the private option key); reviewer answers are NOT copied."""
import re
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
def answers(n):
    ws=load_workbook(f"for_reviewers/review_Reviewer_{n}.xlsx")["Review"]
    return {r[0].value:set(re.findall(r"[A-H]",str(r[11].value or "").upper())) for r in ws.iter_rows(min_row=2)}
A,B=answers("A"),answers("B")
disputed={no for no in A if not (A[no]&B[no])}
wb=load_workbook("for_reviewers/review_Reviewer_A.xlsx")
ws=wb["Review"]
for row in range(ws.max_row,1,-1):
    if ws.cell(row,1).value not in disputed: ws.delete_rows(row)
    else:
        for col in (12,13,14): ws.cell(row,col).value=None
ws.data_validations.dataValidation=[]
dv=DataValidation(type="list",formula1='"A,B,C,D,E,F,G,H,OTHER,NONE"',allow_blank=True)
ws.add_data_validation(dv); dv.add(f"L2:L{ws.max_row}")
ins=wb["Instructions"]; ins["A1"]="Career title mapping review (tie-break set: 28 titles)"
wb.properties.lastModifiedBy="generator"; wb.properties.creator="generator"
wb.save("for_reviewers/review_Reviewer_C_tiebreak.xlsx")
print("disputed titles:",len(disputed))
